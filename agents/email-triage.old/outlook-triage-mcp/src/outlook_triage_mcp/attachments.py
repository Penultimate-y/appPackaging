"""Attachment download and local text extraction (PDF / DOCX / XLSX).

Parsing happens entirely locally — no external service, no LLM tokens.
Extractor libraries are optional: when one is missing, the attachment is
listed with a note instead of failing the whole request.
"""

from __future__ import annotations

import base64
import io
import logging
from typing import Any

log = logging.getLogger(__name__)

try:
    import pypdf

    HAS_PYPDF = True
except ImportError:  # pragma: no cover - environment dependent
    HAS_PYPDF = False

try:
    import docx  # python-docx

    HAS_DOCX = True
except ImportError:  # pragma: no cover - environment dependent
    HAS_DOCX = False

try:
    from openpyxl import load_workbook

    HAS_XLSX = True
except ImportError:  # pragma: no cover - environment dependent
    HAS_XLSX = False

MAX_PDF_PAGES = 50
MAX_XLSX_SHEETS = 5
MAX_XLSX_ROWS = 100


def extract_pdf(data: bytes) -> str:
    if not HAS_PYPDF:
        return "[pypdf not installed — install outlook-triage-mcp[attachments]]"
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
        combined = "\n\n".join(parts).strip()
        return combined or "[PDF has no extractable text — likely scanned/image-only]"
    except Exception as exc:
        return f"[error extracting PDF: {exc}]"


def extract_docx(data: bytes) -> str:
    if not HAS_DOCX:
        return "[python-docx not installed — install outlook-triage-mcp[attachments]]"
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip() or "[DOCX has no extractable text]"
    except Exception as exc:
        return f"[error extracting DOCX: {exc}]"


def extract_xlsx(data: bytes) -> str:
    if not HAS_XLSX:
        return "[openpyxl not installed — install outlook-triage-mcp[attachments]]"
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet_name in workbook.sheetnames[:MAX_XLSX_SHEETS]:
            sheet = workbook[sheet_name]
            parts.append(f"### Sheet: {sheet_name}")
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                if row_count >= MAX_XLSX_ROWS:
                    parts.append("... (additional rows omitted)")
                    break
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    parts.append(" | ".join(cells))
                    row_count += 1
        return "\n".join(parts).strip() or "[XLSX is empty]"
    except Exception as exc:
        return f"[error extracting XLSX: {exc}]"


def extract_text(filename: str, content_type: str, data: bytes) -> str:
    """Dispatch to the right extractor based on filename/content type."""
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    if name.endswith(".pdf") or "pdf" in ctype:
        return extract_pdf(data)
    if name.endswith(".docx") or "wordprocessingml" in ctype:
        return extract_docx(data)
    if name.endswith(".xlsx") or "spreadsheetml" in ctype:
        return extract_xlsx(data)
    if name.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
        return f"[image attachment ({len(data) / 1024:.1f} KB) — not parsed]"
    return f"[unsupported type '{ctype or 'unknown'}' ({len(data) / 1024:.1f} KB) — not parsed]"


def fetch_attachment_texts(
    graph: Any, message_id: str, max_bytes: int = 5 * 1024 * 1024
) -> list[dict[str, Any]]:
    """Download a message's attachments and extract text where possible.

    Returns a list of ``{"name", "size", "content_type", "text"}`` dicts.
    Oversized and non-file attachments are listed with a note, not parsed.
    """
    response = graph.get(
        f"me/messages/{message_id}/attachments", params={"$top": 50}
    )
    results: list[dict[str, Any]] = []

    for attachment in response.get("value", []):
        name = attachment.get("name") or "(unnamed)"
        size = attachment.get("size", 0)
        content_type = (attachment.get("contentType") or "").lower()
        entry = {"name": name, "size": size, "content_type": content_type}

        if attachment.get("@odata.type") != "#microsoft.graph.fileAttachment":
            entry["text"] = "[non-file attachment (e.g. embedded item) — not parsed]"
        elif size > max_bytes:
            entry["text"] = (
                f"[attachment too large ({size / 1024 / 1024:.1f} MB) — not parsed]"
            )
        else:
            try:
                data = base64.b64decode(attachment.get("contentBytes", ""))
            except (ValueError, TypeError) as exc:
                entry["text"] = f"[error decoding attachment content: {exc}]"
            else:
                entry["text"] = extract_text(name, content_type, data)
        results.append(entry)

    return results
